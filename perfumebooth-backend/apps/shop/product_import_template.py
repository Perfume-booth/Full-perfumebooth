from io import BytesIO

import openpyxl
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill, Protection
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation


TEMPLATE_VERSION = '9.0'
MAX_DATA_ROW = 1001
HEADERS = [
    'Handle', 'SKU', 'Title', 'Price', 'Compare At Price', 'Stock', 'Published',
    'Brand', 'Category', 'Product Type', 'Gender', 'Tags', 'SEO Title',
    'SEO Description', 'Description', 'Ideal For', 'Top Notes', 'Middle Notes',
    'Base Notes', 'Size', 'Olfactory Family', 'Fragrance', 'Scent', 'Season',
    'Occasion', 'Dispenser Type', 'Product Form', 'Short Description',
    'Cost Per Item', 'Weight', 'Barcode', 'Material', 'Color', 'Validation Errors',
]
REQUIRED_HEADERS = {'Handle', 'SKU'}
VARIANT_HEADERS = [
    'Product Handle', 'Action', 'Variant SKU', 'Title', 'Option1 Name', 'Option1 Value',
    'Option2 Name', 'Option2 Value', 'Price', 'Compare At Price', 'Stock', 'Status',
    'Validation Errors',
]
IMAGE_HEADERS = [
    'Product Handle', 'Variant SKU', 'Action', 'Image URL', 'Image ID',
    'Alt Text', 'Position', 'Primary', 'Validation Errors',
]
BULK_ACTION_HEADERS = [
    'Target Type', 'Identifier', 'Action', 'Value', 'Field', 'Find', 'Replace', 'Validation Errors',
]


def _style_table_sheet(sheet, include_examples=False, required_headers=None, mode='UPDATE_ONLY'):
    dark = PatternFill('solid', fgColor='26231F')
    gold = PatternFill('solid', fgColor='C89B3C')
    pale_gold = PatternFill('solid', fgColor='FFF2CC')
    red = PatternFill('solid', fgColor='FCE8E6')
    for column, header in enumerate(HEADERS, start=1):
        cell = sheet.cell(1, column, header)
        required_headers = required_headers or REQUIRED_HEADERS
        cell.fill = gold if header in required_headers else dark
        cell.font = Font(color='FFFFFF' if header not in required_headers else '1F1F1F', bold=True)
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.protection = Protection(locked=True)
    sheet.row_dimensions[1].height = 34
    sheet.freeze_panes = 'A2'
    data_end = 3 if include_examples else MAX_DATA_ROW
    validation_column = len(HEADERS)
    validation_letter = get_column_letter(validation_column)
    sheet.auto_filter.ref = f'A1:{validation_letter}{data_end}'
    widths = [22, 20, 34, 14, 18, 12, 14, 22, 22, 22, 16, 32, 32, 42,
              55, 24, 32, 32, 32, 16, 24, 24, 24, 20, 20, 22, 22, 22, 36,
              16, 14, 22, 20, 20, 52]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    for row in sheet.iter_rows(min_row=2, max_row=data_end, min_col=1, max_col=validation_column - 1):
        for cell in row:
            cell.protection = Protection(locked=False)
    if not include_examples:
        for row_number in range(2, MAX_DATA_ROW + 1):
            identity_formula = (
                f'IF(AND(B{row_number}="",C{row_number}="",D{row_number}=""),"SKU, Title and Price required for new products","")'
                if mode in {'CREATE_ONLY', 'UPSERT'} else
                f'IF(AND(A{row_number}="",B{row_number}=""),"Handle or SKU required","")'
            )
            sheet.cell(row_number, validation_column).value = (
                f'=TEXTJOIN("; ",TRUE,{identity_formula},'
                f'IF(AND(D{row_number}<>"",NOT(ISNUMBER(D{row_number}))),"Price must be numeric",""),'
                f'IF(AND(F{row_number}<>"",OR(NOT(ISNUMBER(F{row_number})),F{row_number}<0,MOD(F{row_number},1)<>0)),"Stock must be a non-negative whole number",""))'
            )
            sheet.cell(row_number, validation_column).protection = Protection(locked=True)
        sheet.conditional_formatting.add(
            f'{validation_letter}2:{validation_letter}{MAX_DATA_ROW}',
            FormulaRule(formula=[f'LEN({validation_letter}2)>0'], fill=red),
        )
        last_editable = get_column_letter(validation_column - 1)
        sheet.conditional_formatting.add(f'A2:B{MAX_DATA_ROW}', FormulaRule(formula=[f'AND($A2="",$B2="",COUNTA($C2:${last_editable}2)>0)'], fill=pale_gold))
    sheet.protection.sheet = True
    sheet.protection.password = 'catalog'
    sheet.protection.autoFilter = False
    sheet.protection.selectLockedCells = False
    sheet.protection.selectUnlockedCells = True
    if include_examples:
        sheet.protection.sheet = False


def _add_list_validation(sheet, column, values):
    validation = DataValidation(type='list', formula1=f'"{",".join(values)}"', allow_blank=True)
    validation.error = 'Choose a value from the dropdown.'
    validation.errorTitle = 'Invalid value'
    validation.prompt = 'Select an approved catalog value.'
    validation.promptTitle = 'Catalog value'
    validation.showErrorMessage = True
    validation.showInputMessage = True
    sheet.add_data_validation(validation)
    validation.add(f'{column}2:{column}{MAX_DATA_ROW}')


def _build_variants_sheet(sheet):
    dark = PatternFill('solid', fgColor='26231F')
    gold = PatternFill('solid', fgColor='C89B3C')
    for column, header in enumerate(VARIANT_HEADERS, start=1):
        cell = sheet.cell(1, column, header)
        required = header in {'Product Handle', 'Action', 'Variant SKU'}
        cell.fill = gold if required else dark
        cell.font = Font(color='1F1F1F' if required else 'FFFFFF', bold=True)
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.protection = Protection(locked=True)
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = f'A1:M{MAX_DATA_ROW}'
    for index, width in enumerate([24, 14, 22, 28, 18, 20, 18, 20, 14, 18, 12, 14, 48], start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    for row_number in range(2, MAX_DATA_ROW + 1):
        for cell in sheet[row_number][:12]:
            cell.protection = Protection(locked=False)
        sheet.cell(row_number, 13).value = (
            f'=TEXTJOIN("; ",TRUE,IF(AND(COUNTA(A{row_number}:L{row_number})>0,A{row_number}=""),"Product Handle required",""),'
            f'IF(AND(COUNTA(A{row_number}:L{row_number})>0,B{row_number}=""),"Action required",""),'
            f'IF(AND(COUNTA(A{row_number}:L{row_number})>0,C{row_number}=""),"Variant SKU required",""),'
            f'IF(AND(B{row_number}="ADD",F{row_number}=""),"Option1 Value required for ADD",""),'
            f'IF(AND(B{row_number}="ADD",I{row_number}=""),"Price required for ADD",""))'
        )
        sheet.cell(row_number, 13).protection = Protection(locked=True)
        sheet.cell(row_number, 9).number_format = '0.00'
        sheet.cell(row_number, 10).number_format = '0.00'
        sheet.cell(row_number, 11).number_format = '0'
    _add_list_validation(sheet, 'B', ['ADD', 'UPDATE', 'DELETE'])
    _add_list_validation(sheet, 'L', ['active', 'draft'])
    sheet.conditional_formatting.add(f'M2:M{MAX_DATA_ROW}', FormulaRule(formula=['LEN(M2)>0'], fill=PatternFill('solid', fgColor='FCE8E6')))
    price_validation = DataValidation(type='decimal', operator='greaterThanOrEqual', formula1='0', allow_blank=True)
    price_validation.error = 'Enter a number greater than or equal to zero.'
    price_validation.showErrorMessage = True
    sheet.add_data_validation(price_validation)
    price_validation.add(f'I2:I{MAX_DATA_ROW}')
    compare_validation = DataValidation(
        type='custom', formula1='OR(J2="",J2="__CLEAR__",AND(ISNUMBER(J2),J2>=0))', allow_blank=True
    )
    compare_validation.error = 'Enter a non-negative number, __CLEAR__, or leave blank.'
    compare_validation.showErrorMessage = True
    sheet.add_data_validation(compare_validation)
    compare_validation.add(f'J2:J{MAX_DATA_ROW}')
    stock_validation = DataValidation(type='whole', operator='greaterThanOrEqual', formula1='0', allow_blank=True)
    stock_validation.error = 'Enter a whole number greater than or equal to zero.'
    stock_validation.showErrorMessage = True
    sheet.add_data_validation(stock_validation)
    stock_validation.add(f'K2:K{MAX_DATA_ROW}')
    sheet.protection.sheet = True
    sheet.protection.password = 'catalog'
    sheet.protection.autoFilter = False
    sheet.protection.selectLockedCells = False
    sheet.protection.selectUnlockedCells = True


def _build_images_sheet(sheet):
    dark = PatternFill('solid', fgColor='26231F')
    gold = PatternFill('solid', fgColor='C89B3C')
    for column, header in enumerate(IMAGE_HEADERS, start=1):
        cell = sheet.cell(1, column, header)
        required = header in {'Product Handle', 'Action'}
        cell.fill = gold if required else dark
        cell.font = Font(color='1F1F1F' if required else 'FFFFFF', bold=True)
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.protection = Protection(locked=True)
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = f'A1:I{MAX_DATA_ROW}'
    for index, width in enumerate([24, 22, 14, 60, 38, 40, 12, 12, 48], start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    for row_number in range(2, MAX_DATA_ROW + 1):
        for cell in sheet[row_number][:8]:
            cell.protection = Protection(locked=False)
        sheet.cell(row_number, 9).value = (
            f'=TEXTJOIN("; ",TRUE,IF(AND(COUNTA(A{row_number}:H{row_number})>0,A{row_number}=""),"Product Handle required",""),'
            f'IF(AND(COUNTA(A{row_number}:H{row_number})>0,C{row_number}=""),"Action required",""),'
            f'IF(AND(OR(C{row_number}="ADD",C{row_number}="REPLACE"),D{row_number}=""),"Image URL required",""),'
            f'IF(AND(C{row_number}="DELETE",B{row_number}="",E{row_number}=""),"Image ID or Variant SKU required for DELETE",""))'
        )
        sheet.cell(row_number, 9).protection = Protection(locked=True)
    _add_list_validation(sheet, 'C', ['ADD', 'REPLACE', 'DELETE'])
    _add_list_validation(sheet, 'H', ['yes', 'no'])
    position_validation = DataValidation(type='whole', operator='greaterThanOrEqual', formula1='0', allow_blank=True)
    position_validation.error = 'Enter a whole number greater than or equal to zero.'
    position_validation.showErrorMessage = True
    sheet.add_data_validation(position_validation)
    position_validation.add(f'G2:G{MAX_DATA_ROW}')
    sheet.conditional_formatting.add(f'I2:I{MAX_DATA_ROW}', FormulaRule(formula=['LEN(I2)>0'], fill=PatternFill('solid', fgColor='FCE8E6')))
    sheet.protection.sheet = True
    sheet.protection.password = 'catalog'
    sheet.protection.autoFilter = False
    sheet.protection.selectLockedCells = False
    sheet.protection.selectUnlockedCells = True


def _build_bulk_actions_sheet(sheet):
    dark = PatternFill('solid', fgColor='26231F')
    gold = PatternFill('solid', fgColor='C89B3C')
    for column, header in enumerate(BULK_ACTION_HEADERS, start=1):
        cell = sheet.cell(1, column, header)
        cell.fill = gold if header in {'Target Type', 'Identifier', 'Action', 'Value'} else dark
        cell.font = Font(color='1F1F1F' if column <= 4 else 'FFFFFF', bold=True)
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.protection = Protection(locked=True)
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = f'A1:H{MAX_DATA_ROW}'
    for index, width in enumerate([18, 28, 28, 48, 24, 32, 32, 52], start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    for row_number in range(2, MAX_DATA_ROW + 1):
        for cell in sheet[row_number][:7]:
            cell.protection = Protection(locked=False)
        sheet.cell(row_number, 8).value = (
            f'=TEXTJOIN("; ",TRUE,IF(AND(COUNTA(A{row_number}:D{row_number})>0,A{row_number}=""),"Target Type required",""),'
            f'IF(AND(COUNTA(A{row_number}:D{row_number})>0,B{row_number}=""),"Identifier required",""),'
            f'IF(AND(COUNTA(A{row_number}:D{row_number})>0,C{row_number}=""),"Action required",""),'
            f'IF(AND(COUNTA(A{row_number}:D{row_number})>0,C{row_number}<>"FIND_REPLACE",D{row_number}=""),"Value required",""),'
            f'IF(AND(C{row_number}="FIND_REPLACE",OR(E{row_number}="",F{row_number}="")),"Field and Find required",""))'
        )
        sheet.cell(row_number, 8).protection = Protection(locked=True)
    _add_list_validation(sheet, 'A', ['PRODUCT', 'VARIANT'])
    _add_list_validation(sheet, 'C', [
        'PRICE_PERCENT', 'PRICE_DELTA', 'STOCK_DELTA', 'SET_STATUS',
        'ADD_TAGS', 'REMOVE_TAGS', 'ADD_COLLECTION', 'REMOVE_COLLECTION',
        'PRICE_SET', 'SET_BRAND', 'SET_CATEGORY', 'SET_TAGS', 'FIND_REPLACE',
    ])
    sheet.conditional_formatting.add(f'H2:H{MAX_DATA_ROW}', FormulaRule(formula=['LEN(H2)>0'], fill=PatternFill('solid', fgColor='FCE8E6')))
    sheet.protection.sheet = True
    sheet.protection.password = 'catalog'
    sheet.protection.autoFilter = False
    sheet.protection.selectLockedCells = False
    sheet.protection.selectUnlockedCells = True


def build_product_import_template(mode='UPDATE_ONLY'):
    mode = str(mode).upper()
    required_headers = {'SKU', 'Title', 'Price'} if mode in {'CREATE_ONLY', 'UPSERT'} else REQUIRED_HEADERS
    workbook = openpyxl.Workbook()
    instructions = workbook.active
    instructions.title = 'Instructions'
    products = workbook.create_sheet('Products')
    variants = workbook.create_sheet('Variants')
    images = workbook.create_sheet('Images')
    bulk_actions = workbook.create_sheet('Bulk Actions')
    examples = workbook.create_sheet('Examples')
    meta = workbook.create_sheet('_Template')
    meta.sheet_state = 'hidden'

    instructions['A1'] = f'PerfumeBooth product import template — {mode}'
    instructions['A1'].font = Font(size=18, bold=True, color='7A2E3A')
    instructions['A3'] = 'Template version'
    instructions['B3'] = TEMPLATE_VERSION
    instructions['A5'] = 'How to use'
    instructions['A5'].font = Font(size=13, bold=True)
    first_step = '1. Enter only existing products on the Products sheet.' if mode == 'UPDATE_ONLY' else '1. Enter new products; UPSERT may also contain existing products.'
    mode_rule = {
        'UPDATE_ONLY': 'UPDATE_ONLY never creates a missing product.',
        'CREATE_ONLY': 'CREATE_ONLY never changes an existing product.',
        'UPSERT': 'UPSERT updates matches and creates products that do not exist.',
    }[mode]
    steps = [
        first_step,
        '2. Existing products match by Handle or SKU. New products require SKU, Title and Price; Handle is generated when blank.',
        '3. Include values only in fields you want to change.',
        '4. Leave a cell blank to keep its current database value.',
        '5. Use __CLEAR__ only when you intentionally want to remove an optional value.',
        '6. Review the Validation Errors column before uploading.',
        '7. Upload in Products → Import, inspect the preview, then confirm Apply.',
        '8. Manage variants on the Variants sheet with explicit ADD, UPDATE, or DELETE actions.',
        '9. Every variant row needs its parent Product Handle and a globally unique Variant SKU.',
        '10. Blank UPDATE cells keep their current values; DELETE is blocked while active reservations exist.',
        '11. Manage product gallery and variant images on the Images sheet using HTTPS image URLs.',
        '12. Product images use Image ID for REPLACE/DELETE; variant images use Variant SKU.',
        '13. Mark one product image as Primary. Position controls gallery order.',
        '14. Use Bulk Actions for price percentages/deltas, stock deltas, status, tags, and collections.',
        '15. PRODUCT Identifier is a Handle; VARIANT Identifier is a Variant SKU. Every result is previewed.',
        mode_rule,
    ]
    for row_number, text in enumerate(steps, start=6):
        instructions.cell(row_number, 1, text)
    instructions['A23'] = 'Editable fields'
    instructions['A23'].font = Font(bold=True)
    instructions['A24'] = 'Products: commerce, description, perfume attributes, SEO, shipping and internal cost data. Variants: options, price, stock and status.'
    instructions['A26'] = 'Important'
    instructions['A26'].font = Font(bold=True, color='C00000')
    instructions['A27'] = 'Use standard headers or select a saved Products-sheet mapping. Yellow headers are required (Value is optional for FIND_REPLACE).'
    instructions['A29'] = 'Phase 7: PRICE_SET sets an absolute price; PRICE_DELTA and STOCK_DELTA accept signed values. PRICE_PERCENT accepts signed percentages.'
    instructions['A30'] = 'SET_BRAND / SET_CATEGORY / SET_TAGS assign values. FIND_REPLACE uses Field, Find, Replace (literal, case-sensitive; blank Replace removes text).'
    instructions['A31'] = 'Schedule a price-only preview in MPS. Optional sale end restores the snapshot unless a later price change conflicts. Times use your browser timezone.'
    instructions['A32'] = 'Undo supports existing product/variant updates and bulk actions, not creates/deletes/image changes. Conflicts or reservations block the entire undo.'
    instructions.column_dimensions['A'].width = 115
    instructions.column_dimensions['B'].width = 18
    instructions.freeze_panes = 'A5'

    products.append(HEADERS)
    variants.append(VARIANT_HEADERS)
    _build_variants_sheet(variants)
    images.append(IMAGE_HEADERS)
    _build_images_sheet(images)
    bulk_actions.append(BULK_ACTION_HEADERS)
    _build_bulk_actions_sheet(bulk_actions)
    examples.append(HEADERS)
    _style_table_sheet(products, required_headers=required_headers, mode=mode)
    _style_table_sheet(examples, include_examples=True, required_headers=required_headers, mode=mode)
    if mode == 'UPDATE_ONLY':
        example_data = [
            {'Handle': 'oud-noir', 'SKU': 'PB-OUD-100', 'Price': 1499, 'Published': 'active',
             'Brand': 'PerfumeBooth', 'Category': 'Perfume', 'Product Type': 'Eau de Parfum',
             'Gender': 'Unisex', 'Tags': 'oud, evening', 'Top Notes': 'Saffron',
             'Base Notes': 'Oud, amber', 'Validation Errors': 'Only populated fields would change.'},
            {'Handle': 'rose-mist', 'SKU': 'PB-ROSE-50', 'Title': 'Rose Mist Intense',
             'Compare At Price': '__CLEAR__', 'Stock': 25, 'Published': 'draft', 'Gender': 'Women',
             'SEO Title': 'Rose Mist perfume', 'SEO Description': 'Floral rose fragrance',
             'Validation Errors': 'Blank cells keep their existing values.'},
        ]
    else:
        example_data = [
            {'SKU': 'PB-NEW-100', 'Title': 'New Oud Perfume', 'Price': 1499, 'Stock': 20,
             'Published': 'draft', 'Brand': 'PerfumeBooth', 'Category': 'Perfume',
             'Product Type': 'Eau de Parfum', 'Gender': 'Unisex', 'Tags': 'oud, evening',
             'Description': 'Rich oud fragrance', 'Size': '100 ml',
             'Validation Errors': 'Handle will be generated as new-oud-perfume.'},
            {'Handle': 'rose-intense', 'SKU': 'PB-ROSE-50', 'Title': 'Rose Intense',
             'Price': 999, 'Compare At Price': 1199, 'Stock': 25, 'Published': 'active',
             'Brand': 'PerfumeBooth', 'Category': 'Perfume', 'Product Type': 'Eau de Parfum',
             'Gender': 'Women', 'Tags': 'rose, floral', 'SEO Title': 'Rose Intense',
             'SEO Description': 'Floral rose fragrance', 'Validation Errors': 'Complete new-product example.'},
        ]
    example_rows = [[data.get(header, '') for header in HEADERS] for data in example_data]
    for row_number, values in enumerate(example_rows, start=2):
        for column_number, value in enumerate(values, start=1):
            examples.cell(row_number, column_number, value)
    examples.freeze_panes = 'A2'
    examples.auto_filter.ref = f'A1:{get_column_letter(len(HEADERS))}3'

    for sheet in (products,):
        _add_list_validation(sheet, 'G', ['active', 'draft'])
        _add_list_validation(sheet, 'I', ['Perfume', 'Attar', 'Body Spray', 'Gift Set', 'Scent Shot'])
        _add_list_validation(sheet, 'J', ['Eau de Parfum', 'Eau de Toilette', 'Parfum', 'Perfume Oil', 'Body Mist'])
        _add_list_validation(sheet, 'K', ['Men', 'Women', 'Unisex', 'Kids'])
        price_validation = DataValidation(type='decimal', operator='greaterThanOrEqual', formula1='0', allow_blank=True)
        price_validation.error = 'Enter a number greater than or equal to zero, or leave blank.'
        price_validation.showErrorMessage = True
        sheet.add_data_validation(price_validation)
        price_validation.add(f'D2:D{MAX_DATA_ROW}')
        compare_validation = DataValidation(
            type='custom', formula1='OR(E2="",E2="__CLEAR__",AND(ISNUMBER(E2),E2>=0))', allow_blank=True
        )
        compare_validation.error = 'Enter a non-negative number, __CLEAR__, or leave blank.'
        compare_validation.showErrorMessage = True
        sheet.add_data_validation(compare_validation)
        compare_validation.add(f'E2:E{MAX_DATA_ROW}')
        stock_validation = DataValidation(type='whole', operator='greaterThanOrEqual', formula1='0', allow_blank=True)
        stock_validation.error = 'Enter a whole number greater than or equal to zero, or leave blank.'
        stock_validation.showErrorMessage = True
        sheet.add_data_validation(stock_validation)
        stock_validation.add(f'F2:F{MAX_DATA_ROW}')
        for row_number in range(2, MAX_DATA_ROW + 1):
            sheet.cell(row_number, 4).number_format = '0.00'
            sheet.cell(row_number, 5).number_format = '0.00'
            sheet.cell(row_number, 6).number_format = '0'

    meta['A1'] = 'template_version'
    meta['B1'] = TEMPLATE_VERSION
    meta['A2'] = 'mode'
    meta['B2'] = mode
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
