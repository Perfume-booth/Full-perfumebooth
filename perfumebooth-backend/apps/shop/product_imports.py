import concurrent.futures
import csv
import hashlib
import io
import uuid
from decimal import Decimal, InvalidOperation
from urllib.parse import urlparse

import openpyxl
import requests
from django.db import transaction
from django.core.files.base import ContentFile
from django.utils import timezone
from django.utils.text import slugify

from apps.inventory.models import InventoryMovement, InventoryReservation
from .models import Collection, Inventory, Product, ProductImage, ProductImportJob, ProductVariant
from .services import CollectionSyncService
from .image_imports import download_catalog_image
from .shopify_imports import is_shopify_export, normalize_shopify_csv
from apps.catalog.services.catalog_service import auto_generate_tags


CLEAR = '__CLEAR__'
HEADER_MAP = {
    'handle': 'slug', 'slug': 'slug', 'url slug': 'slug',
    'sku': 'sku', 'variant sku': 'sku',
    'title': 'title', 'product title': 'title',
    'description': 'description', 'body (html)': 'description',
    'price': 'price', 'variant price': 'price',
    'compare at price': 'compare_at_price', 'variant compare at price': 'compare_at_price',
    'compare_at_price': 'compare_at_price',
    'stock': 'stock', 'stock quantity': 'stock', 'inventory qty': 'stock',
    'variant inventory qty': 'stock',
    'published': 'is_online', 'status': 'is_online', 'is_online': 'is_online',
    'brand': 'brand', 'vendor': 'brand', 'category': 'category',
    'product category': 'category', 'product type': 'product_type', 'type': 'product_type',
    'tags': 'tags', 'seo title': 'seo_title', 'seo description': 'seo_description',
    'gender': 'gender', 'brand name': 'brand',
    'ideal for': 'ideal_for', 'top notes': 'top_notes', 'middle notes': 'middle_notes',
    'base notes': 'base_notes', 'size': 'size', 'olfactory family': 'olfactory_family',
    'fragrance': 'fragrance', 'scent': 'scent', 'season': 'season',
    'occasion': 'occasion', 'dispenser type': 'dispenser_type',
    'product form': 'product_form', 'short description': 'short_description',
    'cost per item': 'cost_per_item', 'weight': 'weight', 'barcode': 'barcode',
    'material': 'material', 'color': 'color',
}
EDITABLE_FIELDS = set(HEADER_MAP.values()) - {'slug', 'sku'}
ATTRIBUTE_KEYS = {
    'brand': 'brand', 'category': 'category', 'product_type': 'productType',
    'tags': 'tags', 'seo_title': 'meta_title', 'seo_description': 'meta_description',
    'gender': 'gender',
    'ideal_for': 'ideal_for', 'top_notes': 'top_notes', 'middle_notes': 'middle_notes',
    'base_notes': 'base_notes', 'size': 'size', 'olfactory_family': 'olfactory_family',
    'fragrance': 'fragrance', 'scent': 'scent', 'season': 'season',
    'occasion': 'occasion', 'dispenser_type': 'dispenser_type',
    'product_form': 'product_form', 'short_description': 'short_description',
    'weight': 'weight', 'barcode': 'barcode',
    'material': 'material', 'color': 'color',
}
CLEARABLE = {'compare_at_price', 'cost_per_item', *ATTRIBUTE_KEYS.keys()}
VARIANT_HEADER_MAP = {
    'product handle': 'product_handle', 'handle': 'product_handle',
    'action': 'action', 'variant sku': 'sku', 'sku': 'sku',
    'title': 'title', 'variant title': 'title',
    'barcode': 'barcode', 'variant barcode': 'barcode',
    'option1 name': 'option1_name', 'option1 value': 'option1_value',
    'option2 name': 'option2_name', 'option2 value': 'option2_value',
    'price': 'price', 'variant price': 'price',
    'compare at price': 'compare_at_price', 'variant compare at price': 'compare_at_price',
    'stock': 'stock', 'variant stock': 'stock', 'status': 'is_active',
}
VARIANT_EDITABLE = {
    'title', 'option1_name', 'option1_value', 'option2_name', 'option2_value',
    'price', 'compare_at_price', 'barcode', 'stock', 'is_active',
}
VARIANT_CLEARABLE = {'compare_at_price', 'barcode', 'option2_name', 'option2_value'}
IMAGE_HEADER_MAP = {
    'product handle': 'product_handle', 'handle': 'product_handle',
    'variant sku': 'variant_sku', 'action': 'action', 'image url': 'url',
    'url': 'url', 'image id': 'image_id', 'alt text': 'alt_text',
    'position': 'position', 'primary': 'is_primary',
}
BULK_HEADER_MAP = {
    'target type': 'target_type', 'identifier': 'identifier',
    'action': 'action', 'value': 'value',
    'field': 'field', 'find': 'find', 'replace': 'replace',
}


def _cell(value):
    if value is None:
        return ''
    return str(value).strip()


def read_rows(file_obj):
    raw = file_obj.read()
    file_obj.seek(0)
    filename = file_obj.name.lower()
    if filename.endswith('.csv'):
        text = raw.decode('utf-8-sig')
        reader = csv.DictReader(io.StringIO(text))
        headers = reader.fieldnames or []
        rows = list(reader)
        variant_headers, variant_rows, image_headers, image_rows = [], [], [], []
        bulk_headers, bulk_rows = [], []
        if is_shopify_export(headers):
            headers, rows, variant_headers, variant_rows, image_headers, image_rows = normalize_shopify_csv(rows)
    elif filename.endswith('.xlsx'):
        workbook = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        sheet = workbook['Products'] if 'Products' in workbook.sheetnames else workbook.active
        iterator = sheet.iter_rows(values_only=True)
        headers = [_cell(v) for v in next(iterator, ())]
        rows = [dict(zip(headers, values)) for values in iterator]
        variant_headers, variant_rows = [], []
        if 'Variants' in workbook.sheetnames:
            variant_iterator = workbook['Variants'].iter_rows(values_only=True)
            variant_headers = [_cell(v) for v in next(variant_iterator, ())]
            variant_rows = [dict(zip(variant_headers, values)) for values in variant_iterator]
        image_headers, image_rows = [], []
        if 'Images' in workbook.sheetnames:
            image_iterator = workbook['Images'].iter_rows(values_only=True)
            image_headers = [_cell(v) for v in next(image_iterator, ())]
            image_rows = [dict(zip(image_headers, values)) for values in image_iterator]
        bulk_headers, bulk_rows = [], []
        if 'Bulk Actions' in workbook.sheetnames:
            bulk_iterator = workbook['Bulk Actions'].iter_rows(values_only=True)
            bulk_headers = [_cell(v) for v in next(bulk_iterator, ())]
            bulk_rows = [dict(zip(bulk_headers, values)) for values in bulk_iterator]
    else:
        raise ValueError('Legacy .xls is not supported. Save the file as .xlsx or .csv.')
    return raw, headers, rows, variant_headers, variant_rows, image_headers, image_rows, bulk_headers, bulk_rows


def _normalise_bulk_headers(headers):
    if not headers:
        return {}, []
    mapped, errors = {}, []
    for original in headers:
        canonical = BULK_HEADER_MAP.get(_cell(original).lower())
        if canonical:
            if canonical in mapped:
                errors.append({'sheet': 'Bulk Actions', 'row': 1, 'column': original, 'message': f'Duplicate column for {canonical}.'})
            mapped[canonical] = original
    for field, label in (('target_type', 'Target Type'), ('identifier', 'Identifier'), ('action', 'Action'), ('value', 'Value')):
        if field not in mapped:
            errors.append({'sheet': 'Bulk Actions', 'row': 1, 'column': label, 'message': f'{label} column is required.'})
    return mapped, errors


def _preview_bulk_actions(rows, mapped, product_changes, variant_operations, errors, invalid_rows):
    operations, seen = [], set()
    product_conflicts = {(item['product_id'], item['field']) for item in product_changes}
    variant_conflicts = {
        (item.get('variant_id'), field) for item in variant_operations if item['action'] == 'UPDATE'
        for field in item.get('changes', {})
    }
    allowed = {
        'PRICE_PERCENT', 'PRICE_DELTA', 'STOCK_DELTA', 'SET_STATUS',
        'ADD_TAGS', 'REMOVE_TAGS', 'ADD_COLLECTION', 'REMOVE_COLLECTION',
        'SET_BRAND', 'SET_CATEGORY', 'SET_TAGS', 'FIND_REPLACE', 'PRICE_SET',
    }
    for row_number, row in enumerate(rows, start=2):
        values = {field: _cell(row.get(original)) for field, original in mapped.items()}
        for field in ('find', 'replace'):
            if field in mapped:
                raw = row.get(mapped[field])
                values[field] = '' if raw is None else str(raw)
        if not any(values.values()):
            continue
        target_type = values.get('target_type', '').upper()
        identifier = values.get('identifier', '')
        action = values.get('action', '').upper()
        value = values.get('value', '')
        row_error = False
        if target_type not in {'PRODUCT', 'VARIANT'}:
            errors.append({'sheet': 'Bulk Actions', 'row': row_number, 'column': 'Target Type', 'message': 'Use PRODUCT or VARIANT.'}); row_error = True
        if action not in allowed:
            errors.append({'sheet': 'Bulk Actions', 'row': row_number, 'column': 'Action', 'message': f'Unsupported bulk action: {action or "blank"}.'}); row_error = True
        target = Product.objects.filter(slug=slugify(identifier)).first() if target_type == 'PRODUCT' else ProductVariant.objects.filter(sku=identifier).select_related('product').first()
        if not target:
            errors.append({'sheet': 'Bulk Actions', 'row': row_number, 'column': 'Identifier', 'message': 'Target product handle or variant SKU was not found.'}); row_error = True
        if target_type == 'VARIANT' and action in {'ADD_TAGS', 'REMOVE_TAGS', 'SET_TAGS', 'SET_BRAND', 'SET_CATEGORY', 'FIND_REPLACE', 'ADD_COLLECTION', 'REMOVE_COLLECTION'}:
            errors.append({'sheet': 'Bulk Actions', 'row': row_number, 'column': 'Action', 'message': 'Tags and collections are product-only actions.'}); row_error = True
        logical_field = 'price' if action.startswith('PRICE_') else 'stock' if action == 'STOCK_DELTA' else ('is_active' if target_type == 'VARIANT' else 'is_online') if action == 'SET_STATUS' else 'tags' if action.endswith('_TAGS') else f'collection:{slugify(value)}'
        if action in {'SET_BRAND', 'SET_CATEGORY'}:
            logical_field = action[4:].lower()
        if action == 'FIND_REPLACE':
            logical_field = HEADER_MAP.get(values.get('field', '').lower(), '')
            if logical_field not in {'title', 'description', 'brand', 'category', 'seo_title', 'seo_description', 'short_description'}:
                errors.append({'sheet': 'Bulk Actions', 'row': row_number, 'column': 'Field', 'message': 'Choose Title, Description, Brand, Category, SEO Title, SEO Description or Short Description.'}); row_error = True
        target_id = str(target.pk) if target else ''
        dedupe_key = (target_type, target_id, logical_field)
        if dedupe_key in seen:
            errors.append({'sheet': 'Bulk Actions', 'row': row_number, 'column': 'Action', 'message': 'Only one bulk action per target field is allowed.'}); row_error = True
        seen.add(dedupe_key)
        if target and ((target_type == 'PRODUCT' and (target_id, logical_field) in product_conflicts) or
                       (target_type == 'VARIANT' and (target_id, logical_field) in variant_conflicts)):
            errors.append({'sheet': 'Bulk Actions', 'row': row_number, 'column': 'Action', 'message': 'This field is already changed on the Products or Variants sheet.'}); row_error = True
        old_value = new_value = None
        collection = None
        try:
            if target and action in {'PRICE_PERCENT', 'PRICE_DELTA', 'PRICE_SET'}:
                amount = Decimal(value)
                if not amount.is_finite():
                    raise ValueError('Price adjustment must be finite.')
                old_decimal = target.price
                new_decimal = amount if action == 'PRICE_SET' else old_decimal * (Decimal('1') + amount / Decimal('100')) if action == 'PRICE_PERCENT' else old_decimal + amount
                new_decimal = new_decimal.quantize(Decimal('0.01'))
                if not new_decimal.is_finite() or not 0 <= new_decimal <= Decimal('99999999.99'):
                    raise ValueError('Resulting price must be between 0 and 99999999.99.')
                old_value, new_value = str(old_decimal), str(new_decimal)
            elif target and action == 'STOCK_DELTA':
                delta = int(value)
                stock = target.inventory.stock_level if target_type == 'PRODUCT' else target.stock_level
                reserved = target.inventory.reserved_stock if target_type == 'PRODUCT' else target.reserved_stock
                if not reserved <= stock + delta <= 2147483647:
                    raise ValueError('Resulting stock cannot be below reserved stock or exceed a 32-bit integer.')
                old_value, new_value = stock, stock + delta
            elif target and action == 'SET_STATUS':
                new_value = _parse('is_online', value)
                old_value = target.is_online if target_type == 'PRODUCT' else target.is_active
            elif target and action in {'SET_BRAND', 'SET_CATEGORY'}:
                if not value:
                    raise ValueError('Assignment value is required.')
                old_value, new_value = _current(target, logical_field), _parse(logical_field, value)
            elif target and action == 'FIND_REPLACE' and not row_error:
                find = values.get('find', '')
                if not find:
                    raise ValueError('Find cannot be blank. Replacement is literal and case-sensitive.')
                old_value = _current(target, logical_field)
                new_value = _parse(logical_field, (old_value or '').replace(find, values.get('replace', '')))
            elif target and action in {'ADD_TAGS', 'REMOVE_TAGS', 'SET_TAGS'}:
                requested = [item.strip() for item in value.split(',') if item.strip()]
                if not requested:
                    raise ValueError('Enter at least one tag.')
                old_value = list((target.attributes or {}).get('tags') or [])
                if action == 'SET_TAGS':
                    new_value = list(dict.fromkeys(requested))
                elif action == 'ADD_TAGS':
                    new_value = old_value + [item for item in requested if item not in old_value]
                else:
                    remove = {item.lower() for item in requested}
                    new_value = [item for item in old_value if item.lower() not in remove]
            elif target and action in {'ADD_COLLECTION', 'REMOVE_COLLECTION'}:
                collection = Collection.objects.filter(slug=slugify(value), match_type='manual').first()
                if not collection:
                    raise ValueError('Manual collection was not found.')
                old_value = collection.products.filter(pk=target.pk).exists()
                new_value = action == 'ADD_COLLECTION'
        except (InvalidOperation, ValueError, TypeError, Inventory.DoesNotExist) as exc:
            message = str(exc) or 'Enter a valid value for this action.'
            errors.append({'sheet': 'Bulk Actions', 'row': row_number, 'column': 'Value', 'message': message}); row_error = True
        if row_error:
            invalid_rows.add(f'Bulk Actions:{row_number}')
            continue
        if old_value == new_value:
            continue
        operations.append({
            'sheet': 'Bulk Actions', 'row': row_number, 'target_type': target_type,
            'target_id': target_id, 'identifier': identifier, 'action': action,
            'field': logical_field, 'old_value': old_value, 'new_value': new_value,
            'collection_id': str(collection.pk) if collection else '',
        })
    return operations


def _bulk_current(operation, target):
    field = operation['field']
    if field == 'price':
        return str(target.price)
    if field == 'stock':
        return target.inventory.stock_level if operation['target_type'] == 'PRODUCT' else target.stock_level
    if field in {'is_online', 'is_active'}:
        return target.is_online if operation['target_type'] == 'PRODUCT' else target.is_active
    if field in ATTRIBUTE_KEYS or field in {'title', 'description'}:
        return _current(target, field)
    collection = Collection.objects.get(pk=operation['collection_id'])
    return collection.products.filter(pk=target.pk).exists()


def _normalise_image_headers(headers):
    if not headers:
        return {}, []
    mapped, errors = {}, []
    for original in headers:
        canonical = IMAGE_HEADER_MAP.get(_cell(original).lower())
        if canonical:
            if canonical in mapped:
                errors.append({'sheet': 'Images', 'row': 1, 'column': original, 'message': f'Duplicate column for {canonical}.'})
            mapped[canonical] = original
    for field, label in (('product_handle', 'Product Handle'), ('action', 'Action')):
        if field not in mapped:
            errors.append({'sheet': 'Images', 'row': 1, 'column': label, 'message': f'{label} column is required.'})
    return mapped, errors


def _preview_images(rows, mapped, product_creates, variant_operations, errors, invalid_rows):
    operations = []
    planned_handles = {item['slug'] for item in product_creates}
    planned_variant_skus = {item['sku'] for item in variant_operations if item['action'] == 'ADD'}
    seen_targets = set()
    for row_number, row in enumerate(rows, start=2):
        values = {field: _cell(row.get(original)) for field, original in mapped.items()}
        if not any(values.values()):
            continue
        handle = slugify(values.get('product_handle', ''))
        action = values.get('action', '').upper()
        variant_sku = values.get('variant_sku', '')
        image_id = values.get('image_id', '')
        url = values.get('url', '')
        row_error = False
        if not handle:
            errors.append({'sheet': 'Images', 'row': row_number, 'column': 'Product Handle', 'message': 'Product Handle is required.'}); row_error = True
        product = Product.objects.filter(slug=handle).first() if handle else None
        if handle and not product and handle not in planned_handles:
            errors.append({'sheet': 'Images', 'row': row_number, 'column': 'Product Handle', 'message': 'No existing or newly-created product matches this handle.'}); row_error = True
        if action not in {'ADD', 'REPLACE', 'DELETE'}:
            errors.append({'sheet': 'Images', 'row': row_number, 'column': 'Action', 'message': 'Use ADD, REPLACE, or DELETE.'}); row_error = True
        if action in {'ADD', 'REPLACE'}:
            parsed = urlparse(url)
            if parsed.scheme != 'https' or not parsed.hostname or parsed.username or parsed.password:
                errors.append({'sheet': 'Images', 'row': row_number, 'column': 'Image URL', 'message': 'Enter a public HTTPS image URL without credentials.'}); row_error = True
        if variant_sku:
            variant = ProductVariant.objects.filter(sku=variant_sku).select_related('product').first()
            if not variant and variant_sku not in planned_variant_skus:
                errors.append({'sheet': 'Images', 'row': row_number, 'column': 'Variant SKU', 'message': 'No existing or newly-created variant matches this SKU.'}); row_error = True
            elif variant and variant.product.slug != handle:
                errors.append({'sheet': 'Images', 'row': row_number, 'column': 'Variant SKU', 'message': 'Variant belongs to a different product.'}); row_error = True
            target_key = f'variant:{variant_sku}'
            if action == 'ADD':
                action = 'REPLACE'
        else:
            image = None
            if action in {'REPLACE', 'DELETE'}:
                try:
                    uuid.UUID(image_id)
                    image = ProductImage.objects.select_related('product').filter(pk=image_id).first()
                except (ValueError, TypeError, AttributeError):
                    image = None
                if not image or image.product.slug != handle:
                    errors.append({'sheet': 'Images', 'row': row_number, 'column': 'Image ID', 'message': 'A valid Image ID belonging to this product is required.'}); row_error = True
            target_key = f'image:{image_id}' if image_id else f'new:{handle}:{row_number}'
        if target_key in seen_targets:
            errors.append({'sheet': 'Images', 'row': row_number, 'column': 'Image ID / Variant SKU', 'message': 'The same image target appears more than once.'}); row_error = True
        seen_targets.add(target_key)
        try:
            position = int(values.get('position') or 0)
            if position < 0:
                raise ValueError
        except ValueError:
            errors.append({'sheet': 'Images', 'row': row_number, 'column': 'Position', 'message': 'Position must be a non-negative whole number.'}); row_error = True
            position = 0
        primary_value = values.get('is_primary', '').lower()
        if primary_value not in {'', 'yes', 'no', 'true', 'false', '1', '0'}:
            errors.append({'sheet': 'Images', 'row': row_number, 'column': 'Primary', 'message': 'Use yes or no.'}); row_error = True
        if row_error:
            invalid_rows.add(f'Images:{row_number}')
            continue
        operation = {
            'sheet': 'Images', 'row': row_number, 'action': action, 'product_handle': handle,
            'variant_sku': variant_sku, 'image_id': image_id, 'url': url,
            'alt_text': values.get('alt_text', '')[:500], 'position': position,
            'is_primary': primary_value in {'yes', 'true', '1'},
        }
        if not variant_sku and action in {'REPLACE', 'DELETE'}:
            operation['old_updated_at'] = image.updated_at.isoformat()
        operations.append(operation)
    return operations


def _normalise_variant_headers(headers):
    if not headers:
        return {}, []
    mapped, errors = {}, []
    for original in headers:
        canonical = VARIANT_HEADER_MAP.get(_cell(original).lower())
        if canonical:
            if canonical in mapped:
                errors.append({'sheet': 'Variants', 'row': 1, 'column': original, 'message': f'Duplicate column for {canonical}.'})
            mapped[canonical] = original
    for field, label in (('product_handle', 'Product Handle'), ('action', 'Action'), ('sku', 'Variant SKU')):
        if field not in mapped:
            errors.append({'sheet': 'Variants', 'row': 1, 'column': label, 'message': f'{label} column is required.'})
    return mapped, errors


def _parse_variant(field, value):
    if value == CLEAR:
        if field not in VARIANT_CLEARABLE:
            raise ValueError('This variant field cannot be cleared.')
        return None
    if field in {'price', 'compare_at_price', 'cost_per_item'}:
        return _parse(field, value)
    if field == 'stock':
        return _parse(field, value)
    if field == 'is_active':
        return _parse('is_online', value)
    if not value and field in {'title', 'option1_name', 'option1_value'}:
        raise ValueError('This field cannot be empty.')
    if len(value) > 200:
        raise ValueError('Value is too long.')
    return value


def _variant_current(variant, field):
    value = getattr(variant, 'stock_level' if field == 'stock' else field)
    if field in {'price', 'compare_at_price'} and value is not None:
        return str(value)
    return value


def _preview_variants(rows, mapped, product_creates, errors, invalid_rows):
    operations, seen_skus = [], set()
    planned_handles = {item['slug'] for item in product_creates}
    planned_product_skus = {item['sku'].lower() for item in product_creates}
    for row_number, row in enumerate(rows, start=2):
        values = {field: _cell(row.get(original)) for field, original in mapped.items()}
        if not any(values.values()):
            continue
        handle = slugify(values.get('product_handle', ''))
        action = values.get('action', '').upper()
        sku = values.get('sku', '')
        row_error = False
        for field, label in (('product_handle', 'Product Handle'), ('action', 'Action'), ('sku', 'Variant SKU')):
            if not values.get(field):
                errors.append({'sheet': 'Variants', 'row': row_number, 'column': label, 'message': f'{label} is required.'})
                row_error = True
        if action not in {'ADD', 'UPDATE', 'DELETE'}:
            errors.append({'sheet': 'Variants', 'row': row_number, 'column': 'Action', 'message': 'Use ADD, UPDATE, or DELETE.'})
            row_error = True
        sku_key = sku.lower()
        if sku_key and sku_key in seen_skus:
            errors.append({'sheet': 'Variants', 'row': row_number, 'column': 'Variant SKU', 'message': 'Duplicate Variant SKU in this file.'})
            row_error = True
        seen_skus.add(sku_key)
        product = Product.objects.filter(slug=handle).first() if handle else None
        if handle and not product and handle not in planned_handles:
            errors.append({'sheet': 'Variants', 'row': row_number, 'column': 'Product Handle', 'message': 'No existing or newly-created product matches this handle.'})
            row_error = True
        variant = ProductVariant.objects.filter(sku=sku).select_related('product').first() if sku else None
        if action == 'ADD':
            if variant or Product.objects.filter(sku=sku).exists() or sku_key in planned_product_skus:
                errors.append({'sheet': 'Variants', 'row': row_number, 'column': 'Variant SKU', 'message': 'SKU already belongs to a product or variant.'})
                row_error = True
            for field, label in (('price', 'Price'), ('option1_value', 'Option1 Value')):
                if not values.get(field):
                    errors.append({'sheet': 'Variants', 'row': row_number, 'column': label, 'message': f'{label} is required for ADD.'})
                    row_error = True
            payload = {'sheet': 'Variants', 'row': row_number, 'action': action, 'product_handle': handle, 'sku': sku}
            for field in VARIANT_EDITABLE:
                raw_value = values.get(field, '')
                if raw_value == '':
                    continue
                try:
                    payload[field] = _parse_variant(field, raw_value)
                except ValueError as exc:
                    errors.append({'sheet': 'Variants', 'row': row_number, 'column': mapped.get(field, field), 'message': str(exc)})
                    row_error = True
            payload.setdefault('option1_name', 'Size')
            payload.setdefault('title', payload.get('option1_value', sku))
            payload.setdefault('stock', 0)
            payload.setdefault('is_active', True)
            if not row_error:
                operations.append(payload)
        else:
            if not variant:
                errors.append({'sheet': 'Variants', 'row': row_number, 'column': 'Variant SKU', 'message': f'{action or "This action"} requires an existing variant.'})
                row_error = True
            elif variant.product.slug != handle:
                errors.append({'sheet': 'Variants', 'row': row_number, 'column': 'Product Handle', 'message': 'Variant SKU belongs to a different product.'})
                row_error = True
            if action == 'DELETE' and variant and (
                variant.reserved_stock > 0 or InventoryReservation.objects.filter(variant=variant, status=InventoryReservation.Status.ACTIVE).exists()
            ):
                errors.append({'sheet': 'Variants', 'row': row_number, 'column': 'Action', 'message': 'Cannot delete a variant with active inventory reservations.'})
                row_error = True
            if row_error:
                invalid_rows.add(f'Variants:{row_number}')
                continue
            operation = {'sheet': 'Variants', 'row': row_number, 'action': action, 'product_handle': handle,
                         'sku': sku, 'variant_id': str(variant.pk), 'old_values': {}}
            if action == 'UPDATE':
                operation['changes'] = {}
                for field in VARIANT_EDITABLE:
                    raw_value = values.get(field, '')
                    if raw_value == '':
                        continue
                    try:
                        new_value = _parse_variant(field, raw_value)
                    except ValueError as exc:
                        errors.append({'sheet': 'Variants', 'row': row_number, 'column': mapped.get(field, field), 'message': str(exc)})
                        row_error = True
                        continue
                    old_value = _variant_current(variant, field)
                    operation['old_values'][field] = old_value
                    if old_value != new_value:
                        operation['changes'][field] = new_value
                if not operation['changes']:
                    continue
            else:
                operation['old_values'] = {'sku': variant.sku, 'product_handle': variant.product.slug,
                                           'reserved_stock': variant.reserved_stock}
            if not row_error:
                operations.append(operation)
        if row_error:
            invalid_rows.add(f'Variants:{row_number}')
    return operations


def _normalise_headers(headers, aliases=None):
    mapped, errors = {}, []
    for source in aliases or {}:
        if source not in {_cell(header) for header in headers}:
            errors.append({'row': 1, 'column': source, 'message': 'Saved mapping source column is missing from this file.'})
    for original in headers:
        canonical = (aliases or {}).get(_cell(original)) or HEADER_MAP.get(_cell(original).lower())
        if canonical:
            if canonical in mapped:
                errors.append({'row': 1, 'column': original, 'message': f'Duplicate column for {canonical}.'})
            mapped[canonical] = original
    if 'slug' not in mapped and 'sku' not in mapped:
        errors.append({'row': 1, 'column': 'Handle / SKU', 'message': 'A Handle or SKU column is required.'})
    if not EDITABLE_FIELDS.intersection(mapped):
        errors.append({'row': 1, 'column': '', 'message': 'At least one editable product column is required.'})
    return mapped, errors


def _parse(field, value):
    if value == CLEAR:
        if field not in CLEARABLE:
            raise ValueError('This field cannot be cleared.')
        return None
    if field in {'price', 'compare_at_price'}:
        try:
            parsed = Decimal(value).quantize(Decimal('0.01'))
        except (InvalidOperation, ValueError):
            raise ValueError('Enter a valid decimal number.')
        if not parsed.is_finite() or not 0 <= parsed <= Decimal('99999999.99'):
            raise ValueError('Value must be between 0 and 99999999.99.')
        return str(parsed)
    if field == 'stock':
        try:
            parsed = int(value)
        except (ValueError, TypeError):
            raise ValueError('Enter a whole number.')
        if not 0 <= parsed <= 2147483647:
            raise ValueError('Stock must be between 0 and 2147483647.')
        return parsed
    if field == 'is_online':
        lowered = value.lower()
        if lowered in {'active', 'published', 'true', 'yes', '1'}:
            return True
        if lowered in {'draft', 'archived', 'inactive', 'false', 'no', '0'}:
            return False
        raise ValueError('Use active/draft/archived, published, true/false, yes/no, or 1/0.')
    if field == 'tags':
        return [tag.strip() for tag in value.split(',') if tag.strip()]
    if field == 'gender':
        normalized = value.lower().strip().strip('.…')
        combined = {'men & women', 'men and women', 'men ,women', 'women and men'}
        if normalized in combined or 'unisex' in normalized:
            return 'Unisex'
        if 'female' in normalized or 'women' in normalized:
            return 'Women'
        if 'male' in normalized or normalized == 'men':
            return 'Men'
        if normalized == 'kids':
            return 'Kids'
        if normalized:
            raise ValueError('Use Men, Women, Unisex, or Kids.')
        return ''
    if field == 'category' and ('>' in value or 'health & beauty' in value.lower()):
        from .category_utils import clean_fragrance_category
        return clean_fragrance_category(raw_category=value)
    if field == 'title' and not value:
        raise ValueError('Title cannot be empty.')
    if len(value) > (500 if field in {'title', 'brand', 'category', 'product_type'} else 10000):
        raise ValueError('Value is too long.')
    return value


def _current(product, field):
    if field == 'stock':
        try:
            return product.inventory.stock_level
        except Inventory.DoesNotExist:
            return 0
    if field in ATTRIBUTE_KEYS:
        value = (product.attributes or {}).get(ATTRIBUTE_KEYS[field])
        return [] if field == 'tags' and value is None else value
    value = getattr(product, field)
    return str(value) if field in {'price', 'compare_at_price', 'cost_per_item'} and value is not None else value


VALID_MODES = {'UPDATE_ONLY', 'CREATE_ONLY', 'UPSERT'}


def _creation_payload(values, row_number, mapped, errors, invalid_rows, draft_default):
    required = {'sku': 'SKU', 'title': 'Title', 'price': 'Price'}
    for field, label in required.items():
        if not values.get(field):
            errors.append({'row': row_number, 'column': label, 'message': f'{label} is required when creating a product.'})
            invalid_rows.add(row_number)
    proposed_slug = slugify(values.get('slug') or values.get('title', ''))
    if not proposed_slug:
        errors.append({'row': row_number, 'column': 'Handle', 'message': 'A valid Handle could not be generated.'})
        invalid_rows.add(row_number)
    payload = {'row': row_number, 'slug': proposed_slug, 'sku': values.get('sku', '')}
    for field in EDITABLE_FIELDS.intersection(mapped):
        raw_value = values.get(field, '')
        if raw_value == '':
            continue
        try:
            payload[field] = _parse(field, raw_value)
        except ValueError as exc:
            errors.append({'row': row_number, 'column': mapped[field], 'message': str(exc)})
            invalid_rows.add(row_number)
    payload.setdefault('title', values.get('title', ''))
    payload.setdefault('price', '0.00')
    payload.setdefault('stock', 0)
    payload.setdefault('is_online', not draft_default)
    payload.setdefault('brand', 'PerfumeBooth')
    payload.setdefault('category', 'Perfume')
    payload.setdefault('product_type', 'Eau de Parfum')
    payload.setdefault('tags', [])
    return payload


def create_preview(file_obj, user, mode='UPDATE_ONLY', draft_default=True, job=None, progress_callback=None):
    mode = str(mode or 'UPDATE_ONLY').upper()
    if mode not in VALID_MODES:
        raise ValueError('Mode must be UPDATE_ONLY, CREATE_ONLY, or UPSERT.')
    raw, headers, rows, variant_headers, variant_rows, image_headers, image_rows, bulk_headers, bulk_rows = read_rows(file_obj)
    if progress_callback:
        progress_callback(0, 4)
    mapped, errors = _normalise_headers(headers, job.column_mapping if job else None)
    variant_mapped, variant_header_errors = _normalise_variant_headers(variant_headers)
    errors.extend(variant_header_errors)
    image_mapped, image_header_errors = _normalise_image_headers(image_headers)
    errors.extend(image_header_errors)
    bulk_mapped, bulk_header_errors = _normalise_bulk_headers(bulk_headers)
    errors.extend(bulk_header_errors)
    changes, creates, seen, matched, invalid_rows, unmatched = [], [], set(), set(), set(), 0
    seen_skus, seen_slugs = set(), set()
    nonempty_rows = 0

    for row_number, row in enumerate(rows, start=2):
        values = {field: _cell(row.get(original)) for field, original in mapped.items()}
        if not any(values.values()):
            continue
        nonempty_rows += 1
        handle, sku = values.get('slug', ''), values.get('sku', '')
        normalized_handle = slugify(handle) if handle else ''
        duplicate = False
        if sku and sku.lower() in seen_skus:
            errors.append({'row': row_number, 'column': 'SKU', 'message': 'Duplicate SKU in this file.'})
            invalid_rows.add(row_number); duplicate = True
        if normalized_handle and normalized_handle in seen_slugs:
            errors.append({'row': row_number, 'column': 'Handle', 'message': 'Duplicate Handle in this file.'})
            invalid_rows.add(row_number); duplicate = True
        if sku: seen_skus.add(sku.lower())
        if normalized_handle: seen_slugs.add(normalized_handle)
        if duplicate:
            continue
        by_handle = Product.objects.filter(slug=handle).first() if handle else None
        by_sku = Product.objects.filter(sku=sku).first() if sku else None
        if by_handle and by_sku and by_handle.pk != by_sku.pk:
            errors.append({'row': row_number, 'column': 'Handle / SKU', 'message': 'Handle and SKU identify different products.'})
            invalid_rows.add(row_number); continue
        product = by_handle or by_sku
        if not product:
            if sku and ProductVariant.objects.filter(sku=sku).exists():
                errors.append({'row': row_number, 'column': 'SKU', 'message': 'SKU already belongs to a product variant.'})
                invalid_rows.add(row_number); continue
            if mode == 'UPDATE_ONLY':
                errors.append({'row': row_number, 'column': 'Handle / SKU', 'message': 'No existing product matches this Handle or SKU. UPDATE_ONLY never creates products.'})
                invalid_rows.add(row_number); unmatched += 1; continue
            payload = _creation_payload(values, row_number, mapped, errors, invalid_rows, draft_default)
            if payload['slug'] in seen_slugs and not normalized_handle:
                errors.append({'row': row_number, 'column': 'Handle', 'message': 'Generated Handle duplicates another row.'})
                invalid_rows.add(row_number)
            seen_slugs.add(payload['slug'])
            if Product.objects.filter(slug=payload['slug']).exists() or Product.objects.filter(sku=payload['sku']).exists():
                errors.append({'row': row_number, 'column': 'Handle / SKU', 'message': 'Handle or SKU already exists.'})
                invalid_rows.add(row_number)
            creates.append(payload)
            continue
        if mode == 'CREATE_ONLY':
            errors.append({'row': row_number, 'column': 'Handle / SKU', 'message': 'Product already exists. CREATE_ONLY never updates products.'})
            invalid_rows.add(row_number); continue
        if str(product.pk) in seen:
            errors.append({'row': row_number, 'column': 'Handle / SKU', 'message': 'The same product appears more than once.'})
            invalid_rows.add(row_number); continue
        seen.add(str(product.pk)); matched.add(str(product.pk))
        for field in EDITABLE_FIELDS.intersection(mapped):
            raw_value = values.get(field, '')
            if raw_value == '':
                continue
            try:
                new_value = _parse(field, raw_value)
            except ValueError as exc:
                errors.append({'row': row_number, 'column': mapped[field], 'message': str(exc)})
                invalid_rows.add(row_number); continue
            old_value = _current(product, field)
            if old_value != new_value:
                changes.append({'row': row_number, 'product_id': str(product.pk), 'handle': product.slug,
                                'sku': product.sku, 'field': field, 'old_value': old_value, 'new_value': new_value})

    if progress_callback:
        progress_callback(1, 4)
    variant_operations = _preview_variants(variant_rows, variant_mapped, creates, errors, invalid_rows) if variant_mapped else []
    if progress_callback:
        progress_callback(2, 4)
    image_operations = _preview_images(
        image_rows, image_mapped, creates, variant_operations, errors, invalid_rows
    ) if image_mapped else []
    if progress_callback:
        progress_callback(3, 4)
    bulk_operations = _preview_bulk_actions(
        bulk_rows, bulk_mapped, changes, variant_operations, errors, invalid_rows
    ) if bulk_mapped else []
    if progress_callback:
        progress_callback(4, 4)
    changed_products = len({change['product_id'] for change in changes})
    variant_counts = {action: sum(item['action'] == action for item in variant_operations) for action in ('ADD', 'UPDATE', 'DELETE')}
    summary = {
        'total_rows': nonempty_rows, 'matched_products': len(matched),
        'changed_products': changed_products, 'changed_fields': len(changes),
        'unchanged_rows': max(len(matched) - changed_products, 0),
        'unmatched_rows': unmatched, 'invalid_rows': len(invalid_rows),
        'products_unaffected': max(Product.objects.count() - changed_products, 0),
        'created_products': len(creates), 'updated_products': changed_products,
        'variant_rows': sum(1 for row in variant_rows if any(_cell(value) for value in row.values())),
        'variants_to_add': variant_counts['ADD'], 'variants_to_update': variant_counts['UPDATE'],
        'variants_to_delete': variant_counts['DELETE'],
        'image_rows': sum(1 for row in image_rows if any(_cell(value) for value in row.values())),
        'images_to_add': sum(item['action'] == 'ADD' for item in image_operations),
        'images_to_replace': sum(item['action'] == 'REPLACE' for item in image_operations),
        'images_to_delete': sum(item['action'] == 'DELETE' for item in image_operations),
        'bulk_action_rows': sum(1 for row in bulk_rows if any(_cell(value) for value in row.values())),
        'bulk_actions': len(bulk_operations),
    }
    status = ProductImportJob.Status.INVALID if errors else ProductImportJob.Status.PREVIEW_READY
    values = {
        'created_by': user, 'filename': file_obj.name[:255],
        'file_sha256': hashlib.sha256(raw).hexdigest(), 'status': status, 'mode': mode,
        'summary': summary, 'changes': changes, 'creates': creates,
        'variant_operations': variant_operations, 'image_operations': image_operations,
        'bulk_operations': bulk_operations, 'errors': errors, 'progress_current': 1, 'progress_total': 1,
        'finished_at': timezone.now(), 'error_message': '',
    }
    if job is None:
        return ProductImportJob.objects.create(**values)
    for field, value in values.items():
        setattr(job, field, value)
    job.save(update_fields=list(values))
    return job


def apply_preview(job_id, user, progress_callback=None):
    job = ProductImportJob.objects.get(pk=job_id)
    if job.created_by_id != user.id and not user.is_superuser:
        raise PermissionError('You can only apply your own import preview.')
    if job.status == ProductImportJob.Status.APPLIED:
        raise ValueError('This import was already applied.')
    if job.status not in {ProductImportJob.Status.PREVIEW_READY, ProductImportJob.Status.APPLYING} or job.errors:
        raise ValueError('Fix the preview errors before applying this import.')

    pending_downloads = [
        (index, operation) for index, operation in enumerate(job.image_operations)
        if operation['action'] in {'ADD', 'REPLACE'}
    ]
    total_downloads = len(pending_downloads)
    if progress_callback:
        progress_callback(0, total_downloads + 5)
    downloaded_images = {}
    if total_downloads > 0:
        session = requests.Session()
        from requests.adapters import HTTPAdapter
        adapter = HTTPAdapter(pool_connections=16, pool_maxsize=16, max_retries=1)
        session.mount('https://', adapter)

        def _fetch(item):
            idx, op = item
            return idx, op, download_catalog_image(op['url'], session=session)

        completed_count = 0
        failed_images = []
        with concurrent.futures.ThreadPoolExecutor(max_workers=min(16, total_downloads)) as executor:
            future_to_item = {executor.submit(_fetch, item): item for item in pending_downloads}
            for future in concurrent.futures.as_completed(future_to_item):
                idx, op = future_to_item[future]
                completed_count += 1
                try:
                    idx, op, res = future.result()
                    downloaded_images[idx] = res
                except Exception as exc:
                    failed_images.append({
                        'sheet': 'Images', 'row': op.get('row'),
                        'column': 'Image URL', 'message': str(exc),
                    })
                if progress_callback and (completed_count % 20 == 0 or completed_count == total_downloads):
                    progress_callback(completed_count, total_downloads + 5)

    return _apply_preview_records(job_id, user, downloaded_images, total_downloads, failed_images=failed_images if total_downloads else None, progress_callback=progress_callback)


@transaction.atomic
def _apply_preview_records(job_id, user, downloaded_images, total_downloads, failed_images=None, progress_callback=None, scheduled=False):
    job = ProductImportJob.objects.select_for_update().get(pk=job_id)
    if job.schedule_state and not scheduled:
        raise ValueError('This job is managed by its schedule. Create a new preview to apply manually.')
    if job.created_by_id != user.id and not user.is_superuser:
        raise PermissionError('You can only apply your own import preview.')
    if job.status == ProductImportJob.Status.APPLIED:
        raise ValueError('This import was already applied.')
    if job.status not in {ProductImportJob.Status.PREVIEW_READY, ProductImportJob.Status.APPLYING} or job.errors:
        raise ValueError('Fix the preview errors before applying this import.')
    ids = sorted({change['product_id'] for change in job.changes})
    products = {str(p.pk): p for p in Product.objects.select_for_update().filter(pk__in=ids)}
    inventories = {str(i.product_id): i for i in Inventory.objects.select_for_update().filter(product_id__in=ids)}
    variant_ids = [item.get('variant_id') for item in job.variant_operations if item.get('variant_id')]
    variants = {str(v.pk): v for v in ProductVariant.objects.select_for_update().select_related('product').filter(pk__in=variant_ids)}
    image_ids = [item['image_id'] for item in job.image_operations if item.get('image_id')]
    images = {str(i.pk): i for i in ProductImage.objects.select_for_update().select_related('product').filter(pk__in=image_ids)}
    bulk_product_ids = [item['target_id'] for item in job.bulk_operations if item['target_type'] == 'PRODUCT']
    bulk_variant_ids = [item['target_id'] for item in job.bulk_operations if item['target_type'] == 'VARIANT']
    bulk_products = {str(p.pk): p for p in Product.objects.select_for_update().filter(pk__in=bulk_product_ids)}
    bulk_variants = {str(v.pk): v for v in ProductVariant.objects.select_for_update().filter(pk__in=bulk_variant_ids)}
    bulk_inventory = {str(i.product_id): i for i in Inventory.objects.select_for_update().filter(product_id__in=bulk_product_ids)}
    # Share in-memory instances when sheets touch different fields of one record.
    bulk_products.update({key: value for key, value in products.items() if key in bulk_products})
    bulk_variants.update({key: value for key, value in variants.items() if key in bulk_variants})
    bulk_inventory.update({key: value for key, value in inventories.items() if key in bulk_inventory})
    collection_ids = [item['collection_id'] for item in job.bulk_operations if item.get('collection_id')]
    list(Collection.objects.select_for_update().filter(pk__in=collection_ids))

    for change in job.changes:
        product = products.get(change['product_id'])
        if not product or _current(product, change['field']) != change['old_value']:
            raise ValueError(f"Product {change['handle']} changed after preview. Preview the file again.")
        if change['field'] == 'stock' and change['new_value'] < (inventories[change['product_id']].reserved_stock if change['product_id'] in inventories else 0):
            raise ValueError('Stock cannot be below current reservations.')
    for payload in job.creates:
        if Product.objects.filter(slug=payload['slug']).exists() or Product.objects.filter(sku=payload['sku']).exists():
            raise ValueError(f"Handle or SKU for row {payload['row']} now exists. Preview the file again.")
    for operation in job.variant_operations:
        if operation['action'] == 'ADD':
            if ProductVariant.objects.filter(sku=operation['sku']).exists() or Product.objects.filter(sku=operation['sku']).exists():
                raise ValueError(f"Variant SKU {operation['sku']} now exists. Preview the file again.")
            continue
        variant = variants.get(operation['variant_id'])
        if not variant or variant.sku != operation['sku'] or variant.product.slug != operation['product_handle']:
            raise ValueError(f"Variant {operation['sku']} changed after preview. Preview the file again.")
        if operation['action'] == 'UPDATE':
            if any(_variant_current(variant, field) != old for field, old in operation['old_values'].items()):
                raise ValueError(f"Variant {operation['sku']} changed after preview. Preview the file again.")
            if operation['changes'].get('stock', variant.stock_level) < variant.reserved_stock:
                raise ValueError('Variant stock cannot be below current reservations.')
        elif variant.reserved_stock > 0 or InventoryReservation.objects.select_for_update().filter(
            variant=variant, status=InventoryReservation.Status.ACTIVE
        ).exists():
            raise ValueError(f"Variant {operation['sku']} has an active reservation and cannot be deleted.")
    for operation in job.image_operations:
        if operation.get('image_id'):
            image = images.get(operation['image_id'])
            if not image or image.product.slug != operation['product_handle'] or image.updated_at.isoformat() != operation['old_updated_at']:
                raise ValueError(f"Image {operation['image_id']} changed after preview. Preview the file again.")
    for operation in job.bulk_operations:
        targets = bulk_products if operation['target_type'] == 'PRODUCT' else bulk_variants
        target = targets.get(operation['target_id'])
        if not target or _bulk_current(operation, target) != operation['old_value']:
            raise ValueError(f"Bulk action target {operation['identifier']} changed after preview. Preview again.")
        if operation['field'] == 'stock':
            reserved = bulk_inventory[operation['target_id']].reserved_stock if operation['target_type'] == 'PRODUCT' else target.reserved_stock
            if operation['new_value'] < reserved:
                raise ValueError('Bulk stock cannot be below current reservations.')

    from .advanced_imports import build_snapshot
    job.snapshot = build_snapshot(job)

    if progress_callback:
        progress_callback(total_downloads + 1, total_downloads + 5)

    created_products = {}
    for payload in job.creates:
        attributes = {}
        for field, key in ATTRIBUTE_KEYS.items():
            value = payload.get(field)
            if value is not None and value != '':
                attributes[key] = value
        attributes['tags'] = auto_generate_tags(attributes)
        product = Product.objects.create(
            slug=payload['slug'], sku=payload['sku'], title=payload['title'],
            description=payload.get('description') or '',
            price=Decimal(payload['price']),
            compare_at_price=Decimal(payload['compare_at_price']) if payload.get('compare_at_price') else None,
            cost_per_item=Decimal(payload['cost_per_item']) if payload.get('cost_per_item') else None,
            is_online=payload['is_online'], attributes=attributes,
            created_by=user, updated_by=user,
        )
        inventory = Inventory.objects.create(product=product, stock_level=payload['stock'])
        if payload['stock']:
            InventoryMovement.objects.create(
                movement_type=InventoryMovement.Type.ADJUSTMENT, product=product,
                stock_delta=payload['stock'], stock_after=payload['stock'], reserved_after=0,
                reason=f'Initial stock from catalog import {job.id}',
            )
        CollectionSyncService.syncProductCollections(product)
        created_products[product.slug] = product
    touched = set()
    for change in job.changes:
        product, field, value = products[change['product_id']], change['field'], change['new_value']
        if field == 'stock':
            inventory = inventories.get(str(product.pk))
            if inventory is None:
                inventory = Inventory.objects.create(product=product, stock_level=value)
                inventories[str(product.pk)] = inventory
                old_stock = 0
            else:
                old_stock = inventory.stock_level
                inventory.stock_level = value
                inventory.save(update_fields=['stock_level', 'updated_at'])
            InventoryMovement.objects.create(
                movement_type=InventoryMovement.Type.ADJUSTMENT, product=product,
                stock_delta=value - old_stock, stock_after=value,
                reserved_after=inventory.reserved_stock,
                reason=f'Catalog import {job.id}',
            )
        elif field in ATTRIBUTE_KEYS:
            attributes = dict(product.attributes or {})
            key = ATTRIBUTE_KEYS[field]
            if value is None:
                attributes.pop(key, None)
            else:
                attributes[key] = value
            product.attributes = attributes
            if field == 'barcode':
                product.barcode = value or ''
        elif field in {'price', 'compare_at_price', 'cost_per_item'}:
            setattr(product, field, Decimal(value) if value is not None else None)
        else:
            setattr(product, field, value)
        product.updated_by = user
        touched.add(str(product.pk))
    for product_id in touched:
        product = products[product_id]
        product.save()
        CollectionSyncService.syncProductCollections(product)
    if progress_callback:
        progress_callback(total_downloads + 2, total_downloads + 5)
    existing_handles = {item['product_handle'] for item in job.variant_operations} - set(created_products)
    parent_products = {p.slug: p for p in Product.objects.filter(slug__in=existing_handles)}
    parent_products.update(created_products)
    variant_by_sku = {variant.sku: variant for variant in variants.values()}
    for operation in job.variant_operations:
        action = operation['action']
        if action == 'ADD':
            product = parent_products.get(operation['product_handle'])
            if not product:
                raise ValueError(f"Product {operation['product_handle']} no longer exists. Preview the file again.")
            variant = ProductVariant.objects.create(
                product=product, sku=operation['sku'], title=operation['title'],
                barcode=operation.get('barcode') or '',
                option1_name=operation['option1_name'], option1_value=operation['option1_value'],
                option2_name=operation.get('option2_name') or '', option2_value=operation.get('option2_value') or '',
                price=Decimal(operation['price']),
                compare_at_price=Decimal(operation['compare_at_price']) if operation.get('compare_at_price') else None,
                stock_level=operation['stock'], is_active=operation['is_active'],
            )
            if variant.stock_level:
                InventoryMovement.objects.create(
                    movement_type=InventoryMovement.Type.ADJUSTMENT, product=product, variant=variant,
                    stock_delta=variant.stock_level, stock_after=variant.stock_level, reserved_after=0,
                    reason=f'Initial variant stock from catalog import {job.id}',
                )
            variant_by_sku[variant.sku] = variant
        elif action == 'UPDATE':
            variant = variants[operation['variant_id']]
            old_stock = variant.stock_level
            for field, value in operation['changes'].items():
                target = 'stock_level' if field == 'stock' else field
                if field in {'price', 'compare_at_price'}:
                    value = Decimal(value) if value is not None else None
                setattr(variant, target, value)
            variant.save()
            if 'stock' in operation['changes'] and variant.stock_level != old_stock:
                InventoryMovement.objects.create(
                    movement_type=InventoryMovement.Type.ADJUSTMENT, product=variant.product, variant=variant,
                    stock_delta=variant.stock_level - old_stock, stock_after=variant.stock_level,
                    reserved_after=variant.reserved_stock, reason=f'Catalog import {job.id}',
                )
        else:
            variants[operation['variant_id']].delete()
            variant_by_sku.pop(operation['sku'], None)
    if progress_callback:
        progress_callback(total_downloads + 3, total_downloads + 5)
    image_handles = {item['product_handle'] for item in job.image_operations}
    image_products = {p.slug: p for p in Product.objects.select_for_update().filter(slug__in=image_handles)}
    image_products.update(created_products)
    for index, operation in enumerate(job.image_operations):
        product = image_products.get(operation['product_handle'])
        if not product:
            raise ValueError(f"Product {operation['product_handle']} no longer exists. Preview again.")
        variant_sku = operation.get('variant_sku')
        if variant_sku:
            variant = variant_by_sku.get(variant_sku) or ProductVariant.objects.select_for_update().filter(sku=variant_sku).first()
            if not variant or variant.product_id != product.id:
                raise ValueError(f"Variant {variant_sku} no longer matches this product. Preview again.")
            if operation['action'] == 'DELETE':
                variant.image = None
            else:
                downloaded = downloaded_images.get(index)
                if downloaded:
                    content, filename, _, _ = downloaded
                    variant.image.save(filename, ContentFile(content), save=False)
            variant.save(update_fields=['image', 'updated_at'])
            continue
        if operation['action'] == 'DELETE':
            image = images[operation['image_id']]
            was_primary = image.is_primary
            image.delete()
            if was_primary:
                product.image = None
                product.save(update_fields=['image', 'updated_at'])
            continue
        downloaded = downloaded_images.get(index)
        if not downloaded:
            continue
        content, filename, width, height = downloaded
        if operation['is_primary']:
            ProductImage.objects.filter(product=product, is_primary=True).update(is_primary=False)
        if operation['action'] == 'ADD':
            image = ProductImage(
                product=product, alt_text=operation['alt_text'], position=operation['position'],
                is_primary=operation['is_primary'], source_url=operation['url'], width=width, height=height,
            )
        else:
            image = images[operation['image_id']]
            image.alt_text = operation['alt_text']
            image.position = operation['position']
            image.is_primary = operation['is_primary']
            image.source_url = operation['url']
            image.width, image.height = width, height
        image.image.save(filename, ContentFile(content), save=False)
        image.save()
        if image.is_primary:
            product.image.save(filename, ContentFile(content), save=False)
            product.save(update_fields=['image', 'updated_at'])
    if progress_callback:
        progress_callback(total_downloads + 4, total_downloads + 5)
    for operation in job.bulk_operations:
        target = (bulk_products if operation['target_type'] == 'PRODUCT' else bulk_variants)[operation['target_id']]
        field, new_value = operation['field'], operation['new_value']
        if field == 'price':
            target.price = Decimal(new_value)
            target.save(update_fields=['price', 'updated_at'])
        elif field == 'stock':
            if operation['target_type'] == 'PRODUCT':
                inventory = bulk_inventory[operation['target_id']]
                old_stock = inventory.stock_level
                inventory.stock_level = new_value
                inventory.save(update_fields=['stock_level', 'updated_at'])
                product, variant, reserved = target, None, inventory.reserved_stock
            else:
                old_stock = target.stock_level
                target.stock_level = new_value
                target.save(update_fields=['stock_level', 'updated_at'])
                product, variant, reserved = target.product, target, target.reserved_stock
            InventoryMovement.objects.create(
                movement_type=InventoryMovement.Type.ADJUSTMENT, product=product, variant=variant,
                stock_delta=new_value - old_stock, stock_after=new_value, reserved_after=reserved,
                reason=f'Advanced bulk import {job.id}',
            )
        elif field in {'is_online', 'is_active'}:
            attribute = 'is_online' if operation['target_type'] == 'PRODUCT' else 'is_active'
            setattr(target, attribute, new_value)
            target.save(update_fields=[attribute, 'updated_at'])
        elif field in ATTRIBUTE_KEYS:
            attributes = dict(target.attributes or {})
            if new_value is None:
                attributes.pop(ATTRIBUTE_KEYS[field], None)
            else:
                attributes[ATTRIBUTE_KEYS[field]] = new_value
            target.attributes = attributes
            target.updated_by = user
            target.save(update_fields=['attributes', 'updated_by', 'updated_at'])
            CollectionSyncService.syncProductCollections(target)
        elif field in {'title', 'description'}:
            setattr(target, field, new_value)
            target.updated_by = user
            target.save(update_fields=[field, 'updated_by', 'updated_at'])
            CollectionSyncService.syncProductCollections(target)
        else:
            collection = Collection.objects.get(pk=operation['collection_id'])
            if new_value:
                collection.products.add(target)
            else:
                collection.products.remove(target)
    fields_to_update = ['status', 'applied_at', 'finished_at', 'progress_current', 'progress_total', 'snapshot']
    if failed_images:
        job.errors = list(job.errors or []) + failed_images
        fields_to_update.append('errors')
    job.status = ProductImportJob.Status.APPLIED
    job.applied_at = timezone.now()
    job.finished_at = job.applied_at
    job.progress_current = total_downloads + 5
    job.progress_total = total_downloads + 5
    job.save(update_fields=fields_to_update)
    if progress_callback:
        progress_callback(job.progress_total, job.progress_total)
    return job
